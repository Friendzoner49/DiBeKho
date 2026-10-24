"""Build DiBeKho-du-lieu-goc.xlsx (gitignored: real customer data) from samples/: the reference list of
products and customers to check the app against.  Run:  PYTHONPATH=. uv run python scripts/make_reference_excel.py
"""

from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from dibekho import db
from dibekho.core import catalog
from dibekho.core.text import fmt_qty
from dibekho.excel.catalog_sheet import read_catalog
from dibekho.excel.delivery_note import read_delivery_note

CATALOG = "samples/danh-muc-mau.xlsx"
NOTES = ["samples/phieu-mau-1.xlsx", "samples/phieu-mau-2.xlsx", "samples/phieu-mau-3.xlsx"]
OUT = "DiBeKho-du-lieu-goc.xlsx"

# ---- products: catalogue rows in order -> SP001..SP041 (what 'Nhập danh mục từ Excel' gives on an empty database)
items = read_catalog(CATALOG)
raw = openpyxl.load_workbook(CATALOG).active
merged_top = {}
for rng in raw.merged_cells.ranges:
    for r in range(rng.min_row, rng.max_row + 1):
        merged_top[r] = rng.min_row

# cross-check with the app itself on an empty in-memory database
conn = db.connect(":memory:")
catalog.import_catalog(conn, items)
app_codes = [p["code"] for p in conn.execute("SELECT code FROM products ORDER BY id")]
assert app_codes == [f"SP{i:03d}" for i in range(1, len(items) + 1)], app_codes

notes = [read_delivery_note(p) for p in NOTES]

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
NEW_FILL = PatternFill("solid", fgColor="FFF4CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet(ws, title, note, headers, rows, widths, fills=None):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([note])
    ws["A2"].font = Font(italic=True, color="595959")
    ws.append([])
    ws.append(headers)
    for c in ws[4]:
        c.font, c.fill, c.border = HEAD, FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows):
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if fills and fills[i]:
                c.fill = fills[i]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{openpyxl.utils.get_column_letter(len(headers))}{ws.max_row}"


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Mặt hàng"
rows, fills = [], []
for i, it in enumerate(items, 1):
    top = merged_top.get(it.row, it.row)
    original = (raw.cell(top, 1).value or "").strip()
    rows.append([i, f"SP{i:03d}", it.name_vi, it.name_zh, it.spec, it.variant, it.unit, it.category, "07/2026",
                 f"File danh mục, dòng {it.row}", original])
    fills.append(None)
# lines on the sample notes that the catalogue does not contain (codes if 'Tạo mặt hàng mới' is chosen in this order)
extra = [("phieu-mau-2.xlsx", 13, "vữa chống thấm 20KG – danh mục chỉ có loại 10KG"),
         ("phieu-mau-2.xlsx", 17, "màng tự dính polymer 1,5mm – có thể chính là SP017/SP018, cần bạn xác nhận"),
         ("phieu-mau-2.xlsx", 20, "sơn lót lạnh – không có trong danh mục")]
for k, (fname, row, why) in enumerate(extra, len(items) + 1):
    note = next(n for n, p in zip(notes, NOTES) if p.endswith(fname))
    ln = next(x for x in note.lines if x.row == row)
    rows.append([k, f"SP{k:03d} (dự kiến)", ln.name_vi, ln.name_zh, ln.spec.split("/")[0].strip(), "",
                 ln.spec.split("/")[-1].strip() if "/" in ln.spec else "", "", note.doc_date.strftime("%d/%m/%Y"),
                 f"Phiếu {note.code} ({fname}), dòng {row}", why])
    fills.append(NEW_FILL)
sheet(ws, "DANH MỤC MẶT HÀNG – DỮ LIỆU GỐC",
      f"Lập {datetime.now():%d/%m/%Y %H:%M} từ file danh mục mẫu (41 dòng) và 3 phiếu mẫu. "
      "Mã SP001–SP041 = thứ tự dòng trong file danh mục, đúng như app gán khi 'Nhập danh mục từ Excel' vào database trống. "
      "Dòng vàng: chỉ có trên phiếu, mã dự kiến nếu bấm 'Tạo mặt hàng mới' theo thứ tự này. "
      "Ngày nhập = thời điểm mặt hàng có trong dữ liệu (file danh mục không có ngày cụ thể).",
      ["STT", "Mã SP", "Tên tiếng Việt", "Tên tiếng Trung", "Quy cách", "Loại / màu", "Đơn vị", "Nhóm hàng",
       "Ngày nhập", "Nguồn", "Tên gốc trong file / ghi chú"],
      rows, [5, 16, 42, 30, 14, 20, 8, 24, 11, 22, 60], fills)

ws2 = wb.create_sheet("Khách hàng")
rows = []
customers = {}
for n, path in zip(notes, NOTES):
    c = customers.setdefault(n.phone, {"note": n, "orders": [], "lines": []})
    c["orders"].append(f"{n.code} – {n.doc_date:%d/%m/%Y}")
    c["lines"] += [f"• {ln.name_vi} | {ln.name_zh} | {ln.spec}: {fmt_qty(ln.qty)}" for ln in n.lines]
    if n.skipped_rows:
        c["lines"] += [f"• (bỏ qua, không có số lượng) {s}" for s in n.skipped_rows]
for i, (phone, c) in enumerate(customers.items(), 1):
    n = c["note"]
    total = sum(ln.qty for nn in notes if nn.phone == phone for ln in nn.lines)
    rows.append([i, n.customer_name, phone, n.contact, n.address, len(c["orders"]), "\n".join(c["orders"]),
                 fmt_qty(total), "\n".join(c["lines"])])
sheet(ws2, "DANH SÁCH KHÁCH HÀNG – DỮ LIỆU GỐC",
      "Lấy từ 3 phiếu giao hàng mẫu. Khách được nhận ra theo số điện thoại. Mặt hàng ghi đúng tên trên phiếu "
      "(chưa ghép mã) để đối chiếu với màn hình Chi tiết khách sau khi ghép.",
      ["STT", "Khách hàng", "Điện thoại", "Người liên hệ", "Địa chỉ", "Số đơn", "Mã đơn – ngày giao",
       "Tổng số lượng", "Mặt hàng đã mua (tên trên phiếu | tên Trung | quy cách: SL)"],
      rows, [5, 34, 13, 16, 40, 8, 18, 10, 90])
wb.save(OUT)
print("saved", OUT, len(items) + len(extra), "products,", len(customers), "customers")
