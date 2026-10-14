"""'Dữ liệu gốc': the product catalog and the customer list as one Excel workbook.

Written from Settings ('Xuất dữ liệu gốc'), readable by people, and read back on another computer
('Nhập dữ liệu gốc') to get exactly the same products (same codes) and customers.

Layout (fixed, so the file reads back whatever the interface language was):
  sheet 1 = products, sheet 2 = customers; A1 starts with MARKER; header in row 4; data from row 5.
  products : STT | code | name_vi | name_zh | spec | variant | unit | category | created | min_stock | active
  customers: STT | name | phone | contact | address | orders (read-only, recomputed by the app)
  sheet 3 = stock, like the customer's monthly stock sheet: code | name | category | variant | stock.
            The customer may correct the stock there; importing the file sets the app's stock to it (a stock
            count). Each stock row is tied to a row of sheet 1 (product_row): by code when the code is unique,
            otherwise by position (both sheets are exported in code order), otherwise by name/category/variant.
            Files exported before the code column existed (name | category | variant | stock) work the same way.
"""

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dibekho.core.stock import stock_levels
from dibekho.core.text import clean, parse_number
from dibekho.i18n import t

MARKER = "DIBEKHO"
HEADER_ROW = 4

_HEAD = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F4E78")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class NotMasterData(ValueError):
    pass


@dataclass
class MasterProduct:
    row: int
    code: str
    name_vi: str
    name_zh: str = ""
    spec: str = ""
    variant: str = ""
    unit: str = ""
    category: str = ""
    created_at: str = ""  # 'YYYY-MM-DD HH:MM:SS' or ''
    min_stock: float = 0
    active: bool = True


@dataclass
class MasterStock:
    row: int
    code: str
    name: str
    stock: float
    product_row: int | None = None  # the row of sheet 1 it belongs to; None = not found


@dataclass
class MasterCustomer:
    row: int
    name: str
    phone: str = ""
    contact: str = ""
    address: str = ""


# ---- write -------------------------------------------------------------------

def export(path: Path, conn) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = t("master.sheet_products")
    products = conn.execute("SELECT * FROM products ORDER BY code").fetchall()
    rows = [[i, p["code"], p["name_vi"], p["name_zh"], p["spec"], p["variant"], p["unit"], p["category"],
             _to_datetime(p["created_at"]), p["min_stock"], t("btn.yes") if p["active"] else t("btn.no")]
            for i, p in enumerate(products, 1)]
    _sheet(ws, t("master.title_products"),
           [t("master.col_no"), t("col.code"), t("master.col_name_vi"), t("col.name_zh"), t("col.spec"),
            t("col.variant"), t("col.unit"), t("col.category"), t("master.col_created"), t("col.min_stock"),
            t("settings.active")],
           rows, [5, 10, 42, 30, 14, 20, 8, 24, 12, 12, 9], date_col=9)

    ws2 = wb.create_sheet(t("master.sheet_customers"))
    customers = conn.execute(
        "SELECT c.*, (SELECT COUNT(*) FROM documents d WHERE d.customer_id = c.id AND d.kind = 'delivery' "
        "AND d.status = 'active') AS orders FROM customers c ORDER BY c.name"
    ).fetchall()
    rows = [[i, c["name"], c["phone"], c["contact"], c["address"], c["orders"]] for i, c in enumerate(customers, 1)]
    _sheet(ws2, t("master.title_customers"),
           [t("master.col_no"), t("col.customer"), t("col.phone"), t("col.contact"), t("col.address"),
            t("master.col_orders")],
           rows, [5, 36, 15, 18, 48, 9])

    ws3 = wb.create_sheet(t("master.sheet_stock"))
    levels = stock_levels(conn)
    # by code = the row order of the monthly sheet the catalog was loaded from
    active = conn.execute("SELECT * FROM products WHERE active = 1 ORDER BY code")
    rows = [[p["code"], _stock_name(p), p["category"], p["variant"], levels.get(p["id"], 0)] for p in active]
    _sheet(ws3, t("master.title_stock"),
           [t("col.code"), t("col.product"), t("col.category"), t("col.variant"), t("col.stock")],
           rows, [10, 52, 26, 22, 12])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _stock_name(p) -> str:
    """Like column A of the monthly sheet: Vietnamese name (with packing) on line 1, Chinese name on line 2."""
    return _name_text(p["name_vi"], p["spec"], p["name_zh"])


def _name_text(name_vi, spec, name_zh) -> str:
    first = name_vi
    if spec and spec.casefold() not in first.casefold():
        first += f" ({spec})"
    return f"{first}\n{name_zh}" if name_zh and name_zh != name_vi else first


def _sheet(ws, title, headers, rows, widths, date_col=None, freeze="C5"):
    ws.append([f"{MARKER} – {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([t("master.note", at=datetime.now().strftime("%d/%m/%Y %H:%M"))])
    ws["A2"].font = Font(italic=True, color="595959")
    ws.append([])
    ws.append(headers)
    for c in ws[HEADER_ROW]:
        c.font, c.fill, c.border = _HEAD, _FILL, _BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = _BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if date_col:
            ws.cell(ws.max_row, date_col).number_format = "dd/mm/yyyy"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{max(ws.max_row, HEADER_ROW)}"


def _to_datetime(text: str):
    try:
        return datetime.fromisoformat(text)
    except (TypeError, ValueError):
        return text


# ---- read --------------------------------------------------------------------

def read(path: Path | str) -> tuple[list[MasterProduct], list[MasterCustomer], list[MasterStock]]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # not an Excel file, damaged, locked
        raise NotMasterData(str(exc)) from exc
    try:
        return _read(wb)
    finally:
        wb.close()


def _read(wb):
    sheets = wb.worksheets
    if len(sheets) < 2 or not all(str(s["A1"].value or "").startswith(MARKER) for s in sheets[:2]):
        raise NotMasterData(t("master.not_master"))
    products, customers = [], []
    for r, row in enumerate(sheets[0].iter_rows(min_row=HEADER_ROW + 1, values_only=True), HEADER_ROW + 1):
        row = list(row) + [None] * 11
        name_vi = clean(row[2])
        if not name_vi:
            continue
        products.append(MasterProduct(
            r, clean(row[1]), name_vi, clean(row[3]), clean(row[4]), clean(row[5]), clean(row[6]), clean(row[7]),
            _created(row[8]), _number(row[9]), _active(row[10]),
        ))
    for r, row in enumerate(sheets[1].iter_rows(min_row=HEADER_ROW + 1, values_only=True), HEADER_ROW + 1):
        row = list(row) + [None] * 6
        name, phone = clean(row[1]), clean(row[2])
        if name or phone:
            customers.append(MasterCustomer(r, name, phone, clean(row[3]), clean(row[4])))
    stock = []
    if len(sheets) > 2 and str(sheets[2]["A1"].value or "").startswith(MARKER):
        stock = _read_stock(sheets[2], products)
    return products, customers, stock


def _read_stock(ws, products: list[MasterProduct]) -> list[MasterStock]:
    """Rows of the stock sheet, each tied to a row of sheet 1 (see the module docstring)."""
    header = [clean(c.value) for c in ws[HEADER_ROW]]
    with_code = len([h for h in header if h]) >= 5
    lines = []  # (row, code, name, category, variant, stock)
    for r, row in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True), HEADER_ROW + 1):
        row = list(row) + [None] * 5
        if not with_code:
            row = [None] + row
        code, name = clean(row[0]), clean(row[1])
        if code or name:
            lines.append((r, code, name, clean(row[2]), clean(row[3]), _stock_value(row[4])))

    active = [p for p in products if p.active]
    code_count: dict[str, int] = {}
    for p in products:
        code_count[p.code] = code_count.get(p.code, 0) + 1
    by_code = {p.code: p.row for p in products if p.code and code_count[p.code] == 1}
    by_text: dict[tuple, list[int]] = {}
    for p in active:
        key = (clean(_name_text(p.name_vi, p.spec, p.name_zh)), p.category, p.variant)
        by_text.setdefault(key, []).append(p.row)
    same_order = len(lines) == len(active)

    result = []
    for i, (r, code, name, category, variant, value) in enumerate(lines):
        if code and code in by_code:
            product_row = by_code[code]
        elif same_order:  # exported together: row i of this sheet is the i-th product of sheet 1
            product_row = active[i].row
        else:
            rows = by_text.get((name, category, variant))
            product_row = rows.pop(0) if rows else None
        if value is not None:
            result.append(MasterStock(r, code, name, value, product_row))
    return result


def _stock_value(value) -> float | None:
    if value is None or clean(value) == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return parse_number(clean(value))
    except ValueError:
        return None


def _created(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d 00:00:00")
    text = clean(value)
    try:  # '2026-10-14 09:36:32' typed or kept as text
        return datetime.fromisoformat(text).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).strftime("%Y-%m-%d 00:00:00")
        except ValueError:
            pass
    return ""


def _number(value) -> float:
    try:
        return float(str(value).replace(",", ".")) if value not in (None, "") else 0
    except ValueError:
        return 0


def _active(value) -> bool:
    text = clean(value).casefold()
    return text not in ("không", "khong", "no", "0", "false", "否")  # "No" in any app language
