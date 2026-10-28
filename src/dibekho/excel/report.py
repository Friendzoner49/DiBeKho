"""Excel exports: period statistics (Thống kê) and one customer's purchases with prices (Chi tiết khách).

Money columns are the selling prices read from the delivery notes (no purchase price, no profit).
"""

from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dibekho.core import queries, stats
from dibekho.core.stock import LOCATION_CODES
from dibekho.i18n import t

SHEETS = ("summary", "details", "customers", "top")
_HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
_ORDER_FILL = PatternFill("solid", fgColor="F2F2F2")
_BOLD = Font(bold=True)
MONEY = "#,##0"


def _place(location_id) -> str:
    return t(f"loc.{LOCATION_CODES[location_id]}") if location_id in LOCATION_CODES else ""


def _money_line(m: stats.Money) -> str:
    if m.fees is None:
        return t("stats.money_goods", goods=f"{round(m.goods):,}".replace(",", "."), orders=m.orders)
    return t("stats.money", total=f"{round(m.total):,}".replace(",", "."),
             goods=f"{round(m.goods):,}".replace(",", "."), fees=f"{round(m.fees):,}".replace(",", "."),
             orders=m.orders)


def export(path: Path, conn, start: date, end: date, filters: stats.Filters, by: str,
           sheets: set[str], period_label: str, scope_label: str) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    info = [
        t("report.period", period=period_label),
        t("report.scope", scope=scope_label),
        _money_line(stats.money(conn, start, end, filters)),
        t("report.exported_at", at=datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    if "summary" in sheets:
        _summary(wb, conn, start, end, filters, by, info)
    if "details" in sheets:
        _details(wb, conn, start, end, filters, info)
    if "customers" in sheets:
        _customers(wb, conn, start, end, filters, info)
    if "top" in sheets:
        _top(wb, conn, start, end, filters, info)
    if not wb.worksheets:
        raise ValueError("no sheet selected")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _sheet(wb, title: str, info: list[str], headers: list[str]):
    ws = wb.create_sheet(title[:31])
    for i, line in enumerate(info, start=1):
        ws.cell(row=i, column=1, value=line).font = Font(italic=True, color="555555")
    header_row = len(info) + 2
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font, c.fill = _BOLD, _HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    return ws, header_row + 1


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write(ws, r: int, values: list, money_cols=(), date_cols=()) -> None:
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=col, value=v)
        if col in money_cols:
            cell.number_format = MONEY
        elif col in date_cols:
            cell.number_format = "DD/MM/YYYY"


def _summary(wb, conn, start, end, filters, by, info):
    rows = stats.summary(conn, start, end, filters, by)
    labels = [b.label for b in stats.buckets(start, end, by)] if by != "none" else []
    customer_filtered = bool(filters.customer_ids)
    headers = [t("col.code"), t("col.product"), t("col.name_zh"), t("col.spec"), t("col.variant"), t("col.unit"),
               t("col.category")]
    if not customer_filtered:
        headers.append(t("col.opening"))
    headers += [t("col.in"), t("col.out"), t("col.adjust")]
    if not customer_filtered:
        headers.append(t("col.closing"))
    headers += [t("col.revenue"), t("col.stock_note")]
    money_col = len(headers) - 1
    for label in labels:
        headers += [f"{t('col.in')} {label}", f"{t('col.out')} {label}"]
    ws, r = _sheet(wb, t("sheet.summary"), info, headers)
    for row in rows:
        values = [row.code, row.name_vi, row.name_zh, row.spec, row.variant, row.unit, row.category]
        if not customer_filtered:
            values.append(row.opening)
        values += [row.qty_in, row.qty_out, row.adjust]
        if not customer_filtered:
            values.append(row.closing)
        values += [row.revenue, row.stock_note or None]
        for label in labels:
            qin, qout = row.per_bucket.get(label, (0, 0))
            values += [qin or None, qout or None]
        _write(ws, r, values, money_cols=(money_col,))
        r += 1
    _widths(ws, [9, 45, 20, 16, 14, 8, 14] + [11] * (money_col - 8) + [14, 36] + [11] * (len(headers) - money_col - 1))


def _details(wb, conn, start, end, filters, info):
    kinds = {
        "delivery_out": t("kind.delivery"), "delivery_reverse": t("kind.delivery_cancel"),
        "receipt_in": t("kind.receipt"), "receipt_reverse": t("kind.receipt_cancel"),
        "count_adjust": t("kind.count"), "count_reverse": t("kind.count_cancel"),
        "transfer_out": t("move.transfer_out"), "transfer_in": t("move.transfer_in"),
        "transfer_reverse": t("move.transfer_reverse"),
        "trial_out": t("move.trial_out"), "trial_reverse": t("move.trial_reverse"),
    }
    headers = [t("col.date"), t("col.kind"), t("col.doc_code"), t("col.customer"), t("col.place"), t("col.code"),
               t("col.product"), t("col.spec"), t("col.variant"), t("col.unit"), t("col.qty_change"),
               t("col.unit_price"), t("col.amount"), t("col.note")]
    ws, r = _sheet(wb, t("sheet.details"), info, headers)
    for m in stats.detail_lines(conn, start, end, filters):
        kind = kinds.get(m["kind"], m["kind"])
        if m["receipt_source"] == "return" and m["kind"] in ("receipt_in", "receipt_reverse"):
            kind += " · " + t("receipt_form.source_short_return")
        values = [date.fromisoformat(m["occurred_on"]), kind, m["code"], m["customer_name"], _place(m["location_id"]),
                  m["product_code"], m["name_vi"], m["spec"], m["variant"], m["unit"], m["qty"], m["unit_price"],
                  m["amount"], m["note"]]
        _write(ws, r, values, money_cols=(12, 13), date_cols=(1,))
        r += 1
    _widths(ws, [11, 16, 10, 30, 12, 9, 40, 16, 14, 8, 10, 12, 14, 30])


def _customers(wb, conn, start, end, filters, info):
    headers = [t("col.customer"), t("col.code"), t("col.product"), t("col.spec"), t("col.variant"), t("col.unit"),
               t("col.total_qty"), t("col.revenue"), t("col.times"), t("col.last_date")]
    ws, r = _sheet(wb, t("sheet.customers"), info, headers)
    for row in stats.customer_product(conn, start, end, filters):
        values = [row["customer_name"], row["code"], row["name_vi"], row["spec"], row["variant"], row["unit"],
                  row["total_qty"], row["revenue"], row["times"], date.fromisoformat(row["last_date"])]
        _write(ws, r, values, money_cols=(8,), date_cols=(10,))
        r += 1
    _widths(ws, [32, 9, 40, 16, 14, 8, 12, 14, 10, 12])


def _top(wb, conn, start, end, filters, info):
    headers = [t("col.rank"), t("col.code"), t("col.product"), t("col.spec"), t("col.variant"), t("col.unit"),
               t("col.out"), t("col.revenue")]
    ws, r = _sheet(wb, t("sheet.top"), info, headers)
    for i, row in enumerate(stats.top_out(conn, start, end, filters, limit=50), start=1):
        _write(ws, r, [i, row.code, row.name_vi, row.spec, row.variant, row.unit, row.qty_out, row.revenue],
               money_cols=(8,))
        r += 1
    _widths(ws, [6, 9, 45, 16, 14, 8, 12, 14])


# ---- one customer ---------------------------------------------------------------------------------------------

def export_customer(path: Path, conn, customer_id: int, start: date | None, end: date | None,
                    period_label: str) -> Path:
    """What one customer bought and at what price, in a period (the customer window with its date filter):
    sheet 1 every order with its lines, fees and total; sheet 2 the products with quantity, revenue and
    price range."""
    c = queries.customer(conn, customer_id)
    money = queries.customer_money(conn, customer_id, start, end)
    fmt = lambda v: f"{round(v or 0):,}".replace(",", ".")  # noqa: E731
    info = [
        t("report.customer", name=c["name"], phone=c["phone"] or "—"),
        t("report.period", period=period_label),
        t("cust.money_value", total=fmt(money["total"]), orders=money["orders"], goods=fmt(money["goods"]),
          fees=fmt(money["fees"])),
        t("report.exported_at", at=datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    headers = [t("col.date"), t("col.doc_code"), t("col.code"), t("col.product"), t("col.name_zh"), t("col.spec"),
               t("col.variant"), t("col.qty"), t("col.unit"), t("col.unit_price"), t("col.amount"), t("col.note")]
    ws, r = _sheet(wb, t("sheet.customer_orders"), info, headers)
    lines: dict[int, list] = {}
    for ln in queries.customer_lines(conn, customer_id, start, end):
        lines.setdefault(ln["document_id"], []).append(ln)
    for d in reversed(queries.documents(conn, "delivery", start, end, customer_id=customer_id)):  # oldest first
        day = date.fromisoformat(d["doc_date"])
        for ln in lines.get(d["id"], []):
            _write(ws, r, [day, d["code"], ln["product_code"], ln["name_vi"], ln["name_zh"], ln["spec"], ln["variant"],
                           ln["qty"], ln["unit"], ln["unit_price"], ln["amount"], ln["note"]],
                   money_cols=(10, 11), date_cols=(1,))
            r += 1
        _write(ws, r, [day, d["code"], None, t("report.order_total"), None, None, None, d["total_qty"], None, None,
                       d["goods_amount"]], money_cols=(11,), date_cols=(1,))
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).font = _BOLD
            ws.cell(row=r, column=col).fill = _ORDER_FILL
        r += 1
        if d["fee_amount"]:  # apart, not in the revenue
            _write(ws, r, [day, d["code"], None, t("cust.fees"), None, None, None, None, None, None, d["fee_amount"]],
                   money_cols=(11,), date_cols=(1,))
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).font = Font(italic=True, color="777777")
            r += 1
    _widths(ws, [11, 9, 9, 40, 20, 16, 14, 8, 8, 13, 15, 30])

    headers = [t("col.code"), t("col.product"), t("col.name_zh"), t("col.spec"), t("col.variant"), t("col.total_qty"),
               t("col.unit"), t("col.revenue"), t("report.price_low"), t("report.price_high"), t("col.times"),
               t("col.last_date")]
    ws, r = _sheet(wb, t("sheet.customer_products"), info, headers)
    for p in queries.customer_products(conn, customer_id, start, end):
        _write(ws, r, [p["code"], p["name_vi"], p["name_zh"], p["spec"], p["variant"], p["total_qty"], p["unit"],
                       p["revenue"], p["min_price"], p["max_price"], p["times"], date.fromisoformat(p["last_date"])],
               money_cols=(8, 9, 10), date_cols=(12,))
        r += 1
    _widths(ws, [9, 40, 20, 16, 14, 10, 8, 15, 13, 13, 9, 12])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
