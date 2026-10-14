"""'Dữ liệu gốc': export products and customers to Excel, import them on another computer."""

import openpyxl
import pytest

from dibekho import db
from dibekho.core import catalog
from dibekho.excel import master_data
from tests.test_features import setup_notes

PRODUCT_COLS = "code, name_vi, name_zh, spec, variant, unit, category, min_stock, active, created_at"
CUSTOMER_COLS = "name, contact, phone, phone_norm, address"


def snapshot(conn):
    return ([tuple(r) for r in conn.execute(f"SELECT {PRODUCT_COLS} FROM products ORDER BY code")],
            [tuple(r) for r in conn.execute(f"SELECT {CUSTOMER_COLS} FROM customers ORDER BY name")])


def test_export_then_import_on_an_empty_computer_gives_the_same_data(conn, ws, tmp_path):
    setup_notes(conn, ws)
    with db.transaction(conn):
        pid = catalog.create_product(conn, "Màng PPE", "高分子防水卷材", "16.5KG", category="Màng", min_stock=5,
                                     variant="Nguyên sinh")
        catalog.update_product(conn, pid, active=0)
    path = master_data.export(tmp_path / "goc.xlsx", conn)

    wb = openpyxl.load_workbook(path)
    customers_sheet = wb.worksheets[1]
    assert [c.value for c in customers_sheet[4]] == ["STT", "Khách hàng", "Điện thoại", "Người liên hệ", "Địa chỉ",
                                                     "Số đơn"]
    orders = {r[1]: r[5] for r in customers_sheet.iter_rows(min_row=5, values_only=True)}
    assert orders == {"Khách Một": 2, "Khách Hai": 1}

    other = db.connect(":memory:")
    result = catalog.import_master(other, *master_data.read(path))
    assert (result["products_new"], result["customers_new"], result["skipped"]) == (3, 2, [])
    assert snapshot(other) == snapshot(conn)

    # importing again changes nothing and creates no duplicates
    result = catalog.import_master(other, *master_data.read(path))
    assert (result["products_new"], result["products_updated"], result["customers_new"]) == (0, 3, 0)
    assert snapshot(other) == snapshot(conn)


def test_edited_file_updates_and_never_duplicates(conn, ws, tmp_path):
    setup_notes(conn, ws)
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    wb = openpyxl.load_workbook(path)
    sheet = wb.worksheets[0]
    sheet.cell(5, 10).value = 12  # low-stock alert of the first product
    sheet.append([None, "", "Sơn lót lạnh", "冷底油", "15KG"])  # new product without a code
    wb.save(path)
    result = catalog.import_master(conn, *master_data.read(path))
    assert (result["products_new"], result["products_updated"]) == (1, 2)
    first = conn.execute("SELECT min_stock FROM products ORDER BY code LIMIT 1").fetchone()[0]
    assert first == 12
    assert conn.execute("SELECT code FROM products WHERE name_vi = 'Sơn lót lạnh'").fetchone()[0] == "SP003"


def test_same_product_under_another_code_is_skipped(conn, ws, tmp_path):
    setup_notes(conn, ws)
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    name_vi, name_zh, spec = conn.execute("SELECT name_vi, name_zh, spec FROM products WHERE code = 'SP002'").fetchone()
    other = db.connect(":memory:")
    other.execute("INSERT INTO products(code, name_vi, name_zh, spec) VALUES('SP009', ?, ?, ?)", (name_vi, name_zh, spec))
    result = catalog.import_master(other, *master_data.read(path))
    assert result["products_new"] == 1 and result["skipped"] == [(6, "SP002", "SP009")]
    assert other.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 2  # no duplicate


def test_rejects_other_excel_files(tmp_path):
    from tests.conftest import make_note

    with pytest.raises(master_data.NotMasterData):
        master_data.read(make_note(tmp_path / "phieu.xlsx"))


def test_file_exported_in_chinese_reads_the_same(conn, ws, tmp_path):
    from dibekho import i18n

    setup_notes(conn, ws)
    with db.transaction(conn):
        pid = catalog.create_product(conn, "Màng PPE", spec="16.5KG")
        catalog.update_product(conn, pid, active=0)
    i18n.set_language("zh")
    try:
        path = master_data.export(tmp_path / "goc-zh.xlsx", conn)
    finally:
        i18n.set_language("vi")
    other = db.connect(":memory:")
    catalog.import_master(other, *master_data.read(path))
    assert snapshot(other) == snapshot(conn)  # '否' (inactive) included


def test_export_has_a_stock_sheet_like_the_monthly_sheet(conn, ws, tmp_path):
    from datetime import date

    from dibekho.core.stock import stock_count

    a, b = setup_notes(conn, ws)
    with db.transaction(conn):
        stock_count(conn, date(2026, 10, 10), {a: 50})
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    sheet = openpyxl.load_workbook(path).worksheets[2]
    assert [c.value for c in sheet[4]] == ["Mã", "Mặt hàng", "Nhóm hàng", "Loại / màu", "Tồn"]
    rows = {r[1].splitlines()[1]: r for r in sheet.iter_rows(min_row=5, values_only=True)}
    assert rows["背水涂"][4] == 50 and rows["背水涂"][0].startswith("SP")
    assert rows["固沙剂"][4] == -5  # never counted: 1 + 4 delivered
    other = db.connect(":memory:")
    catalog.import_master(other, *master_data.read(path))
    assert snapshot(other) == snapshot(conn)


def _levels_by_code(conn):
    from dibekho.core.stock import stock_levels

    levels = stock_levels(conn)
    return {r["code"]: levels.get(r["id"], 0) for r in conn.execute("SELECT id, code FROM products")}


def test_stock_corrected_in_the_file_is_imported_exactly(conn, ws, tmp_path):
    from datetime import date

    from dibekho.core.stock import stock_count

    a, b = setup_notes(conn, ws)
    with db.transaction(conn):
        stock_count(conn, date(2026, 10, 10), {a: 50})
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    wb = openpyxl.load_workbook(path)
    sheet = wb.worksheets[2]
    for row in sheet.iter_rows(min_row=5):
        row[4].value = str(len(row[1].value))  # the customer types new numbers (even as text)
    wanted = {row[0].value: float(row[4].value) for row in sheet.iter_rows(min_row=5)}
    wb.save(path)

    # same computer: the stock becomes exactly the file's numbers
    result = catalog.import_master(conn, *master_data.read(path))
    assert result["stock_counted"] == 2 and result["stock_skipped"] == []
    assert _levels_by_code(conn) == wanted
    # another, empty computer: same products, same stock
    other = db.connect(":memory:")
    catalog.import_master(other, *master_data.read(path))
    assert _levels_by_code(other) == wanted


def test_old_stock_sheet_without_code_column_is_matched_by_name(conn, ws, tmp_path):
    setup_notes(conn, ws)
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    wb = openpyxl.load_workbook(path)
    sheet = wb.worksheets[2]
    sheet.delete_cols(1)  # files exported before 14/10/2026: name | category | variant | stock
    sheet["A1"] = "DIBEKHO – BẢNG TỒN KHO"  # the title was in column A
    for row in sheet.iter_rows(min_row=5):
        row[3].value = 7
    wb.save(path)
    products, customers, stock = master_data.read(path)
    assert [st.product_row for st in stock] == [p.row for p in products]  # same order as sheet 1
    other = db.connect(":memory:")
    catalog.import_master(other, products, customers, stock)
    assert set(_levels_by_code(other).values()) == {7}


def test_master_file_in_input_is_recognised(conn, ws):
    from dibekho.core import inbox

    setup_notes(conn, ws)
    master_data.export(ws.input / "goc.xlsx", conn)
    [item] = inbox.items(conn, ws)
    assert item.kind == "master" and len(item.master[2]) == 2


def test_rows_sharing_a_code_stay_separate_products_with_their_stock(conn, ws, tmp_path):
    """The user copied a row for another variant and kept its code: both rows must become products,
    and the stock of each row must go to the right one (the bug that merged PPE variants)."""
    setup_notes(conn, ws)
    path = master_data.export(tmp_path / "goc.xlsx", conn)
    wb = openpyxl.load_workbook(path)
    products_sheet, stock_sheet = wb.worksheets[0], wb.worksheets[2]
    products_sheet.append([3, "SP001", products_sheet["C5"].value, "", products_sheet["E5"].value, "Tái chế",
                           "thùng", "", None, 0, "Có"])  # copy of SP001, another variant, same code
    stock_sheet.append(["SP001", "copy", "", "Tái chế", 9])
    for row, qty in zip(stock_sheet.iter_rows(min_row=5), (11, 22, 9)):
        row[4].value = qty
    wb.save(path)

    other = db.connect(":memory:")
    result = catalog.import_master(other, *master_data.read(path))
    assert result["renumbered"] == [(7, "SP001", "SP003")]
    assert result["products_new"] == 3
    assert _levels_by_code(other) == {"SP001": 11, "SP002": 22, "SP003": 9}


def old_file_with_stock(conn, path, stock_rows):
    """A file exported before the code column existed; the stock sheet holds stock_rows
    (name, category, variant, stock) in the order the user left them."""
    master_data.export(path, conn)
    wb = openpyxl.load_workbook(path)
    sheet = wb.worksheets[2]
    sheet.delete_cols(1)
    sheet["A1"] = "DIBEKHO – BẢNG TỒN KHO"
    sheet.delete_rows(5, sheet.max_row)
    for row in stock_rows:
        sheet.append(list(row))
    wb.save(path)
    return path


def ppe_variants(conn):
    """Like the user's file: same name and packing, told apart only by category and type."""
    with db.transaction(conn):
        white = catalog.create_product(conn, "MÀNG CHỐNG THẤM PPE", "高分子聚乙烯涤纶防水卷材", "18.5KG",
                                       category="纯白色", variant="瑞德牌")  # SP001
        brand = catalog.create_product(conn, "MÀNG CHỐNG THẤM PPE", "高分子聚乙烯涤纶防水卷材", "18.5KG",
                                       category="瑞德牌", variant="瑞德牌")  # SP002
    return white, brand


PPE_NAME = "MÀNG CHỐNG THẤM PPE (18.5KG)\n高分子聚乙烯涤纶防水卷材"


def test_the_stock_sheet_decides_which_product_gets_the_stock(conn, tmp_path):
    """The stock sheet is the reference: its rows are matched by what they say, not by row order
    (the user's file had the two sheets in a different order: by order, 79 and 31 were swapped)."""
    ppe_variants(conn)
    path = old_file_with_stock(conn, tmp_path / "goc.xlsx", [
        (PPE_NAME, "瑞德牌", "瑞德牌", 79),   # row 5 describes SP002
        (PPE_NAME, "瑞德牌", "纯白色", 31),   # row 6 describes SP001 (labels typed in other columns)
    ])
    products, customers, stock = master_data.read(path)
    assert [st.matched_by for st in stock] == ["name", "name"]
    by_row = {p.row: p for p in products}
    assert [by_row[st.product_row].code for st in stock] == ["SP002", "SP001"]
    assert not any(master_data.stock_differs(st, by_row[st.product_row]) for st in stock)
    sp001 = by_row[stock[1].product_row]
    assert master_data.label_change(stock[1], sp001) == ("纯白色 · 瑞德牌", "瑞德牌 · 纯白色")
    other = db.connect(":memory:")
    result = catalog.import_master(other, products, customers, stock)
    assert _levels_by_code(other) == {"SP001": 31, "SP002": 79}
    # the stock sheet's category and type win
    assert result["labels_updated"] == ["SP001"]
    assert tuple(other.execute("SELECT category, variant FROM products WHERE code = 'SP001'").fetchone()) == (
        "瑞德牌", "纯白色")


def test_same_name_other_type_is_another_product_and_takes_the_stock_sheet_labels(conn, tmp_path):
    """The user's rows 49/50: same name, 纯白色 and 再生料 in the stock sheet; the products sheet says
    纯原料·瑞德牌 and 再生料·瑞德牌. 再生料 goes to 再生料, the 纯白色 row to the other product, which
    becomes 纯白色 as the stock sheet says. Not a duplicate, not a problem."""
    with db.transaction(conn):
        for category in ("纯原料", "再生料"):
            catalog.create_product(conn, "MÀNG CHỐNG THẤM PPE", "高分子聚乙烯涤纶防水卷材", "16.5KG",
                                   category=category, variant="瑞德牌")
    name = "MÀNG CHỐNG THẤM PPE (16.5KG)\n高分子聚乙烯涤纶防水卷材"
    path = old_file_with_stock(conn, tmp_path / "goc.xlsx", [(name, "纯白色", "纯白色", 0), (name, "再生料", "再生料", 3)])
    products, customers, stock = master_data.read(path)
    assert master_data.same_rows(stock) == {}
    by_row = {p.row: p for p in products}
    assert [by_row[st.product_row].code for st in stock] == ["SP001", "SP002"]
    assert not any(master_data.stock_differs(st, by_row[st.product_row]) for st in stock)
    catalog.import_master(conn, products, customers, stock)
    rows = conn.execute("SELECT code, category, variant FROM products ORDER BY code").fetchall()
    assert [tuple(r) for r in rows] == [("SP001", "纯白色", "纯白色"), ("SP002", "再生料", "再生料")]
    assert _levels_by_code(conn) == {"SP001": 0, "SP002": 3}


def test_rows_identical_in_the_stock_sheet_are_reported(conn, tmp_path):
    ppe_variants(conn)
    path = old_file_with_stock(conn, tmp_path / "goc.xlsx", [(PPE_NAME, "瑞德牌", "瑞德牌", 79), (PPE_NAME, "瑞德牌", "瑞德牌", 5)])
    _, _, stock = master_data.read(path)
    assert master_data.same_rows(stock) == {5: [6], 6: [5]}


def test_vietnamese_name_edited_in_one_sheet_only(conn, tmp_path):
    """Sheet 1 got 'HY98' added to the Vietnamese name, the stock sheet did not: the Chinese name and the
    packing still find the product."""
    with db.transaction(conn):
        for category in ("再生料", "纯原料"):
            catalog.create_product(conn, "PPE HY98", "高分子聚乙烯涤纶防水卷材 （HY98）", "16.5KG", category=category)
    name = "PPE (16.5KG)\n高分子聚乙烯涤纶防水卷材 （HY98）"
    path = old_file_with_stock(conn, tmp_path / "goc.xlsx", [(name, "纯原料", "", 4), (name, "再生料", "", 6)])
    products, customers, stock = master_data.read(path)
    by_row = {p.row: p for p in products}
    assert [by_row[st.product_row].code for st in stock] == ["SP002", "SP001"]
    assert not any(master_data.stock_differs(st, by_row[st.product_row]) for st in stock)


def test_unknown_stock_rows_fall_back_to_row_order_and_are_flagged(conn, tmp_path):
    ppe_variants(conn)
    path = old_file_with_stock(conn, tmp_path / "goc.xlsx", [("Tên khác hẳn", "", "", 5), ("Tên khác nữa", "", "", 6)])
    products, _, stock = master_data.read(path)
    assert master_data.needs_check(stock)
    assert [st.matched_by for st in stock] == ["position", "position"]
    by_row = {p.row: p for p in products}
    assert all(master_data.stock_differs(st, by_row[st.product_row]) for st in stock)


def test_file_with_codes_needs_no_check(conn, ws, tmp_path):
    setup_notes(conn, ws)
    products, _, stock = master_data.read(master_data.export(tmp_path / "moi.xlsx", conn))
    assert not master_data.needs_check(stock)
    by_row = {p.row: p for p in products}
    assert not any(master_data.stock_differs(st, by_row[st.product_row]) for st in stock)
