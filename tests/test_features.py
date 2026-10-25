import json
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from dibekho import db
from dibekho.core import backup, catalog, deliveries, inbox, queries, stats, transfer
from dibekho.core.stock import close_document, stock_count, stock_levels
from dibekho.core.workspace import Config, Workspace
from dibekho.excel import report
from dibekho.excel.receipt import write_template
from tests.conftest import make_note

A = ("背水涂\nvữa chống thấm", "10KG/thùng")
B = ("固沙剂\nchất cố định cát", "10KG/thùng")


def setup_notes(conn, ws):
    make_note(ws.deliveries / "1.xlsx", code="1", day="02-09-2026", customer="Khách Một", phone="0901",
              lines=[(*A, 3, ""), (*B, 1, "")])
    make_note(ws.deliveries / "2.xlsx", code="2", day="15-09-2026", customer="Khách Hai", phone="0902",
              lines=[(*A, 5, "")])
    make_note(ws.deliveries / "3.xlsx", code="3", day="03-10-2026", customer="Khách Một", phone="0901",
              lines=[(*A, 2, ""), (*B, 4, "giao sau")])
    deliveries.scan(conn, ws)
    for u in deliveries.unknown_lines(conn, ws):
        deliveries.create_product_for_line(conn, u)
    deliveries.scan(conn, ws)
    ids = {r["name_zh"]: r["id"] for r in conn.execute("SELECT id, name_zh FROM products")}
    return ids["背水涂"], ids["固沙剂"]


def test_receipt_template_roundtrip(conn, ws):
    a, b = setup_notes(conn, ws)
    path = write_template(ws.input / "hang-ve.xlsx", catalog.products(conn), date(2026, 9, 1))
    wb = openpyxl.load_workbook(path)
    sheet = wb.active
    for row in sheet.iter_rows(min_row=7):
        if row[0].value == conn.execute("SELECT code FROM products WHERE id=?", (a,)).fetchone()[0]:
            row[5].value = 20
    wb.save(path)

    [item] = inbox.items(conn, ws)
    assert item.kind == "receipt" and item.receipt.doc_date == date(2026, 9, 1)
    inbox.apply_receipt(conn, ws, item)
    assert stock_levels(conn)[a] == 20 - 3 - 5 - 2
    assert not (ws.input / "hang-ve.xlsx").exists()
    assert len(list(ws.processed.iterdir())) == 1

    # the same file dropped again is refused
    processed = next(ws.processed.iterdir())
    (ws.input / "again.xlsx").write_bytes(processed.read_bytes())
    [again] = inbox.items(conn, ws)
    assert again.kind == "error"


def test_delivery_note_in_input_gets_helpful_error(conn, ws):
    make_note(ws.input / "phieu.xlsx")
    [item] = inbox.items(conn, ws)
    assert item.kind == "error" and "phieu-giao-hang" in item.error


def test_stock_count_sets_stock(conn, ws):
    a, b = setup_notes(conn, ws)
    with db.transaction(conn):
        stock_count(conn, date(2026, 8, 31), {a: 50, b: 10})
    levels = stock_levels(conn)
    assert levels[a] == 50 - 3 - 5 - 2
    assert levels[b] == 10 - 1 - 4


def test_cancel_reverses_and_keeps_history(conn, ws):
    a, _ = setup_notes(conn, ws)
    doc = conn.execute("SELECT id FROM documents WHERE code = '2'").fetchone()[0]
    with db.transaction(conn):
        close_document(conn, doc)
    assert stock_levels(conn)[a] == -5
    assert conn.execute("SELECT status FROM documents WHERE id = ?", (doc,)).fetchone()[0] == "cancelled"
    with pytest.raises(ValueError):
        with db.transaction(conn):
            close_document(conn, doc)


def test_period_summary(conn, ws):
    a, b = setup_notes(conn, ws)
    with db.transaction(conn):
        stock_count(conn, date(2026, 8, 31), {a: 50, b: 10})
    s, e = stats.month_range(2026, 9)
    rows = {r.product_id: r for r in stats.summary(conn, s, e)}
    assert (rows[a].opening, rows[a].qty_out, rows[a].closing) == (50, 8, 42)
    assert (rows[b].opening, rows[b].qty_out, rows[b].closing) == (10, 1, 9)

    q = stats.summary(conn, *stats.quarter_range(2026, 3), by="month")
    qa = next(r for r in q if r.product_id == a)
    assert qa.per_bucket["09/2026"] == (0, 8)

    one = stats.summary(conn, s, e, stats.Filters(customer_ids=[_customer(conn, "Khách Hai")]))
    assert [(r.product_id, r.qty_out, r.opening) for r in one] == [(a, 5, None)]

    only_b = stats.summary(conn, s, e, stats.Filters(product_ids=[b]))
    assert [r.product_id for r in only_b] == [b]


def _customer(conn, name):
    return conn.execute("SELECT id FROM customers WHERE name = ?", (name,)).fetchone()[0]


def test_customer_products(conn, ws):
    a, b = setup_notes(conn, ws)
    rows = queries.customer_products(conn, _customer(conn, "Khách Một"))
    got = {r["id"]: (r["total_qty"], r["times"], r["last_date"]) for r in rows}
    assert got == {a: (5, 2, "2026-10-03"), b: (5, 2, "2026-10-03")}


def test_search_accent_insensitive(conn, ws):
    setup_notes(conn, ws)
    res = queries.search(conn, "vua chong tham")
    assert len(res["products"]) == 1
    assert queries.search(conn, "khach mot")["customers"]
    assert queries.search(conn, "背水")["products"]
    assert queries.search(conn, "giao sau")["documents"]
    assert [d["code"] for d in queries.search(conn, "003")["documents"]] == ["3"]


def test_report_export(conn, ws, tmp_path):
    setup_notes(conn, ws)
    path = report.export(tmp_path / "r.xlsx", conn, *stats.quarter_range(2026, 3), stats.Filters(), "month",
                         set(report.SHEETS), "Q3/2026", "Tất cả")
    wb = openpyxl.load_workbook(path)
    assert len(wb.sheetnames) == 4
    text = " ".join(str(c.value) for ws_ in wb.worksheets for row in ws_.iter_rows() for c in row if c.value)
    assert "VND" not in text  # selling prices are in their own columns, never in the text
    assert "Doanh thu" in text and "Đơn giá" in text and "Thành tiền" in text


def test_transfer_roundtrip(tmp_path):
    ws_a = Workspace(tmp_path / "A")
    ws_a.ensure()
    conn_a = db.connect(ws_a.db_path)
    setup_notes(conn_a, ws_a)
    cfg = Config(workspace=str(ws_a.root), device_id="a1", device_name="MayA")
    pkg = transfer.export_package(conn_a, ws_a, cfg)
    assert not transfer.has_unsent_changes(conn_a)
    levels_a = stock_levels(conn_a)

    ws_b = Workspace(tmp_path / "B")
    ws_b.ensure()
    conn_b = db.connect(ws_b.db_path)
    with db.transaction(conn_b):
        catalog.create_product(conn_b, "Hàng của máy B")
    assert transfer.has_unsent_changes(conn_b)
    target = ws_b.input / pkg.name
    target.write_bytes(pkg.read_bytes())
    [item] = inbox.items(conn_b, ws_b)
    assert item.kind == "package" and item.manifest.device_name == "MayA"

    backup.named_backup(conn_b, ws_b, "truoc-nhap-du-lieu")
    conn_b.close()
    manifest = transfer.import_package(ws_b, target)
    conn_b = db.connect(ws_b.db_path)
    transfer.after_import(conn_b, manifest)
    assert stock_levels(conn_b) == levels_a
    assert len(list(ws_b.originals.glob("*.xlsx"))) == 3
    # B's own folder has none of A's files: no 'missing' warnings
    assert deliveries.scan(conn_b, ws_b).pending == 0
    # B later gets a copy of one of A's notes: recognised, not counted twice
    (ws_b.deliveries / "1.xlsx").write_bytes((ws_a.deliveries_done / "1.xlsx").read_bytes())
    assert deliveries.scan(conn_b, ws_b).imported == []
    assert stock_levels(conn_b) == levels_a


def test_daily_backup_keeps_10(conn, ws):
    for day in range(1, 20):
        backup.daily_backup(conn, ws, date(2026, 9, day))
    assert len(list(ws.backups.glob("DiBeKho_*.sqlite"))) == backup.KEEP
    assert backup.daily_backup(conn, ws, date(2026, 9, 19)) is None


def test_backup_older_than_30_days_is_pruned(conn, ws, monkeypatch):
    # well under KEEP (10): only the age rule can explain the removal
    old = backup.snapshot(conn, ws.backups / "DiBeKho_2026-01-01.sqlite")
    recent = [backup.snapshot(conn, ws.backups / f"DiBeKho_2026-09-{d:02d}.sqlite") for d in range(15, 20)]
    old_time = (datetime.now() - timedelta(days=31)).timestamp()
    os.utime(old, (old_time, old_time))
    backup._prune(ws)
    assert not old.exists()
    assert all(p.exists() for p in recent)


def test_backup_floor_keeps_newest_3_even_if_old(conn, ws):
    paths = [backup.snapshot(conn, ws.backups / f"DiBeKho_2026-01-{d:02d}.sqlite") for d in range(1, 4)]
    old_time = (datetime.now() - timedelta(days=60)).timestamp()
    for p in paths:
        os.utime(p, (old_time, old_time))
    backup._prune(ws)
    assert all(p.exists() for p in paths)


def test_maintenance_runs_every_two_weeks(conn, ws):
    cfg = Config.create(ws.root)
    assert backup.maintenance_due(cfg)
    backup.run_maintenance(conn, ws, cfg)
    assert cfg.last_cleanup == date.today().isoformat()
    assert not backup.maintenance_due(cfg)


def test_every_translation_key_exists():
    root = Path(__file__).resolve().parent.parent / "src" / "dibekho"
    vi = json.loads((root / "i18n" / "vi.json").read_text(encoding="utf-8"))
    for lang in ("en", "zh"):
        other = json.loads((root / "i18n" / f"{lang}.json").read_text(encoding="utf-8"))
        assert set(vi) == set(other), (lang, set(vi) ^ set(other))
        for key, text in vi.items():  # same {placeholders}, or t() fails at run time
            assert set(re.findall(r"\{(\w+)\}", text)) == set(re.findall(r"\{(\w+)\}", other[key])), (lang, key)
    used = set()
    for f in root.rglob("*.py"):
        used |= set(re.findall(r'\bt\(\s*"([a-z_.0-9]+)"', f.read_text(encoding="utf-8")))
    used |= {f"search.{k}" for k in ("products", "customers", "documents")}
    dynamic = (
        [f"status.{s}" for s in ("active", "cancelled", "replaced")]
        + [f"kind.{k}" for k in ("delivery", "receipt", "stock_count")]
        + [f"kind.{k}_plural" for k in ("delivery", "receipt", "stock_count")]
        + [f"move.{k}" for k in ("delivery_out", "delivery_reverse", "receipt_in", "receipt_reverse",
                                 "count_adjust", "count_reverse")]
        + [f"file.{k}" for k in ("error", "changed", "missing")]
        + [f"inbox.{k}" for k in ("receipt", "catalog", "package", "error")]
        + [f"period.{k}" for k in ("day", "month", "quarter", "year", "range")]
        + [f"split.{k}" for k in ("none", "day", "month")]
        + [f"sheet.{k}" for k in report.SHEETS]
    )
    missing = (used | set(dynamic)) - set(vi) - {"key"}
    assert not missing, missing


def test_chinese_input(conn, ws):
    from dibekho.core.text import parse_number

    setup_notes(conn, ws)
    assert [p["name_zh"] for p in queries.search(conn, "背水")["products"]] == ["背水涂"]
    assert queries.search(conn, "固沙剂")["products"][0]["name_zh"] == "固沙剂"
    assert parse_number("１２") == 12 and parse_number("１２，５") == 12.5 and parse_number(" 3,5 ") == 3.5


def test_catalog_sheet_dropped_into_input_is_recognised(conn, ws):
    from tests.test_catalog import make_sheet

    make_sheet(ws.input / "danh-muc.xlsx")
    [item] = inbox.items(conn, ws)
    assert item.kind == "catalog" and len(item.catalog) == 6


def test_unusable_file_in_input_is_hidden_not_moved(conn, ws):
    path = ws.input / "ghi-chu.xlsx"
    openpyxl.Workbook().save(path)
    [item] = inbox.items(conn, ws)
    assert item.kind == "error"
    inbox.hide(conn, item)
    assert inbox.items(conn, ws) == [] and path.exists()
    assert not ws.processed.exists() or not any(ws.processed.iterdir())


def test_discarded_receipt_can_still_be_imported_from_da_xu_ly(conn, ws):
    from dibekho.excel.receipt import write_template

    a, _ = setup_notes(conn, ws)
    path = write_template(ws.input / "hang-ve.xlsx", catalog.products(conn), date(2026, 9, 1)) or ws.input / "hang-ve.xlsx"
    wb = openpyxl.load_workbook(path)
    for row in wb.active.iter_rows(min_row=7):
        if row[0].value == conn.execute("SELECT code FROM products WHERE id=?", (a,)).fetchone()[0]:
            row[5].value = 7
    wb.save(path)
    [item] = inbox.items(conn, ws)
    inbox.discard(ws, item)  # 'Bỏ' in Cần xử lý
    [moved] = list(ws.processed.iterdir())
    again = inbox.read_item(conn, moved)  # 'Chọn file để nhập…' on the goods-in page
    assert again.kind == "receipt"
    inbox.apply_receipt(conn, ws, again)
    assert moved.exists()  # a file picked outside input/ is left where it is
    assert stock_levels(conn)[a] == 7 - 3 - 5 - 2


def test_waiting_notes_from_the_other_computer_do_not_stay_in_to_do(tmp_path):
    """Machine A exports while a note still waits for product names; on B that file does not exist."""
    ws_a = Workspace(tmp_path / "A")
    ws_a.ensure()
    conn_a = db.connect(ws_a.db_path)
    setup_notes(conn_a, ws_a)
    make_note(ws_a.deliveries / "moi.xlsx", code="99", lines=[("新产品\nHàng mới tinh", "5KG/thùng", 1, "")])
    deliveries.scan(conn_a, ws_a)
    assert deliveries.pending_count(conn_a) == 1
    pkg = transfer.export_package(conn_a, ws_a, Config(workspace=str(ws_a.root), device_id="a", device_name="A"))

    ws_b = Workspace(tmp_path / "B")
    ws_b.ensure()
    db.connect(ws_b.db_path).close()
    manifest = transfer.import_package(ws_b, pkg)
    conn_b = db.connect(ws_b.db_path)
    transfer.after_import(conn_b, manifest)
    assert deliveries.scan(conn_b, ws_b).pending == 0
    assert inbox.items(conn_b, ws_b) == []

    # the note copied to B later: asked again
    (ws_b.deliveries / "moi.xlsx").write_bytes((ws_a.deliveries / "moi.xlsx").read_bytes())
    deliveries.scan(conn_b, ws_b)
    assert [u.name_vi for u in deliveries.unknown_lines(conn_b, ws_b)] == ["Hàng mới tinh"]


def test_delete_product_only_when_unused(conn, ws):
    a, _ = setup_notes(conn, ws)
    spare = catalog.create_product(conn, "Sơn thừa", spec="5KG")
    catalog.set_alias(conn, "son thua|5kg", spare, "Sơn thừa")
    assert catalog.product_documents(conn, a) == 3
    assert catalog.product_documents(conn, spare) == 0
    with pytest.raises(ValueError):
        catalog.delete_product(conn, a)
    catalog.delete_product(conn, spare)
    assert conn.execute("SELECT COUNT(*) FROM products WHERE id = ?", (spare,)).fetchone()[0] == 0
    assert catalog.alias_product(conn, "son thua|5kg") is None


def test_a_deleted_code_is_never_given_again(conn):
    first = catalog.create_product(conn, "A")
    middle = catalog.create_product(conn, "B")
    last = catalog.create_product(conn, "C")
    catalog.delete_product(conn, last)  # SP003, the highest
    catalog.delete_product(conn, middle)  # SP002, a gap
    new = catalog.create_product(conn, "D")
    codes = [r[0] for r in conn.execute("SELECT code FROM products ORDER BY code")]
    assert codes == ["SP001", "SP004"] and new != first


def test_merge_products_keeps_smaller_code_and_frees_the_larger(conn):
    from dibekho.core.stock import Line, insert_document

    keep = catalog.create_product(conn, "Sơn A", spec="20KG")  # SP001
    drop = catalog.create_product(conn, "Sơn A", spec="20KG")  # SP002, a duplicate
    insert_document(conn, "delivery", date(2026, 10, 1), [Line(drop, 3)], code="1", code_norm="1")
    catalog.set_alias(conn, "son a|20kg", drop, "Sơn A")
    survivor = catalog.merge_products(conn, drop, keep)  # order given does not matter
    assert survivor == keep
    assert conn.execute("SELECT COUNT(*) FROM products WHERE id = ?", (drop,)).fetchone()[0] == 0
    assert conn.execute("SELECT product_id FROM document_lines").fetchone()[0] == keep
    assert conn.execute("SELECT product_id FROM stock_movements").fetchone()[0] == keep
    assert catalog.alias_product(conn, "son a|20kg") == keep
    new = catalog.create_product(conn, "Sơn B")  # the freed SP002 is given out again
    assert conn.execute("SELECT code FROM products WHERE id = ?", (new,)).fetchone()[0] == "SP002"


def test_restore_deleted_product_from_backup(conn, ws):
    pid = catalog.create_product(conn, "Sơn thừa", "涂料", "5KG", unit="thùng", category="Sơn", min_stock=10)
    catalog.update_product(conn, pid, stock_note="Hộp bị lỗi")
    backup.named_backup(conn, ws, "truoc-xoa")
    catalog.delete_product(conn, pid)
    found = catalog.deleted_products(conn, ws)
    assert [d.code for d in found] == ["SP001"]
    data = found[0]
    assert (data.name_vi, data.name_zh, data.spec, data.unit, data.category, data.min_stock, data.stock_note) == (
        "Sơn thừa", "涂料", "5KG", "thùng", "Sơn", 10, "Hộp bị lỗi")
    new_id = catalog.restore_product(conn, data.code, data)
    assert catalog.find_merge_candidates(conn, new_id, data) == ([], 0)  # nothing else to confuse it with
    row = conn.execute("SELECT * FROM products WHERE id = ?", (new_id,)).fetchone()
    assert row["code"] == "SP001" and row["name_vi"] == "Sơn thừa"


def test_find_merge_candidates_via_an_old_alias_key(conn, ws):
    """The deleted product had a learned name (but, unusually, no document yet -- the only way it could be
    deleted) pointing to it; by the time it is restored, that same name has been relearned onto a different
    product. That is the strongest sign of all that the two are the same item, so it is first and counted
    as an alias match -- the caller pre-ticks it."""
    pid = catalog.create_product(conn, "Sơn thừa", spec="5KG")
    catalog.set_alias(conn, "son thua|5kg", pid, "Sơn thừa")
    backup.named_backup(conn, ws, "truoc-xoa")
    catalog.delete_product(conn, pid)
    again = catalog.create_product(conn, "Sơn thừa khác chữ", "甲", spec="5KG")  # name_vi alone would miss this
    catalog.set_alias(conn, "son thua|5kg", again, "Sơn thừa")
    data = catalog.deleted_products(conn, ws)[0]
    assert data.alias_keys == ("son thua|5kg",)
    new_id = catalog.restore_product(conn, data.code, data)
    candidates, alias_matches = catalog.find_merge_candidates(conn, new_id, data)
    assert candidates and candidates[0]["id"] == again and alias_matches == 1


def test_find_merge_candidates_falls_back_to_the_fuzzy_matcher(conn, ws):
    """The usual case: a catalog product nobody ever recognised on a note (no alias) is deleted by mistake,
    then re-created from a note whose own wording does not exactly match the restored product's. Uses the
    same, higher-confidence matcher as the single best guess 'Cần xử lý' pre-fills for an unrecognised
    line (score >= 0.5), not the looser one listing a few options -- this dialog must stay rare."""
    pid = catalog.create_product(conn, "Sơn FS-101", "雨堡垒FS-101防水涂料", "20KG", unit="thùng", variant="Trắng")
    backup.named_backup(conn, ws, "truoc-xoa")
    catalog.delete_product(conn, pid)
    again = catalog.create_product(conn, "Sơn chống thấm FS-101", "雨堡垒FS-101防水涂料（白色）", "20KG/thùng")
    data = catalog.deleted_products(conn, ws)[0]
    assert data.alias_keys == ()  # it never had one
    assert catalog.find_same(conn, data.name_vi, data.spec, data.variant) is None  # the exact match misses too
    new_id = catalog.restore_product(conn, data.code, data)
    candidates, alias_matches = catalog.find_merge_candidates(conn, new_id, data)
    assert alias_matches == 0
    assert any(c["id"] == again for c in candidates)


def test_merge_products_fills_empty_fields_from_the_dropped_product(conn):
    keep = catalog.create_product(conn, "Sơn A", spec="20KG")  # SP001, no notes typed yet
    drop = catalog.create_product(conn, "Sơn A", "涂料A", "20KG", category="Sơn")
    catalog.update_product(conn, drop, stock_note="2 thùng in lệch", min_stock=15)
    catalog.merge_products(conn, keep, drop)
    row = conn.execute("SELECT * FROM products WHERE id = ?", (keep,)).fetchone()
    assert (row["name_zh"], row["category"], row["stock_note"], row["min_stock"]) == ("涂料A", "Sơn", "2 thùng in lệch", 15)


def test_deleted_products_are_found_in_an_undo_snapshot_and_file_is_untouched(conn, ws):
    from dibekho.core import undo

    pid = catalog.create_product(conn, "Sơn riêng")
    point = undo.mark(conn, ws, "delete_product")
    catalog.delete_product(conn, pid)
    before = point.stat().st_mtime
    codes = [d.code for d in catalog.deleted_products(conn, ws)]
    assert codes == ["SP001"]
    assert point.stat().st_mtime == before


def test_config_keeps_ui_scale(tmp_path, monkeypatch):
    from dibekho.core import workspace

    monkeypatch.setattr(workspace, "config_dir", lambda: tmp_path)
    cfg = workspace.Config.create(tmp_path / "w")
    assert cfg.ui_scale == 1.2
    cfg.ui_scale = 1.4
    cfg.save()
    assert workspace.Config.load().ui_scale == 1.4
